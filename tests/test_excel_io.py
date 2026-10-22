from decimal import Decimal
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from freight_calculator.calculator import calculate
from freight_calculator.excel_io import read_workbook, write_results, write_template
from freight_calculator.models import CargoLine, CalculationSettings, FreightCalculationError


class ExcelIoTests(unittest.TestCase):
    def test_results_export_round_trip_preserves_units_and_settings(self):
        lines = [
            CargoLine(
                '=not a formula', 3, Decimal("12.5"), Decimal("8"), Decimal("4"),
                "in", Decimal("1.2"), "lb",
            ),
            CargoLine(
                "Carton", 2, Decimal("30"), Decimal("20"), Decimal("10"),
                "cm", Decimal("3"), "kg",
            ),
        ]
        settings = CalculationSettings(
            Decimal("5000"), "per_piece", Decimal("0.5")
        )
        result = calculate(lines, settings)
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "roundtrip.xlsx"
            write_results(path, lines, result)
            imported_lines, imported_settings = read_workbook(path)
        self.assertEqual(imported_lines, lines)
        self.assertEqual(imported_settings, settings)
        self.assertEqual(result, calculate(imported_lines, imported_settings))

    def test_empty_template_is_valid_but_has_no_cargo_to_calculate(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "template.xlsx"
            write_template(path)
            with self.assertRaisesRegex(FreightCalculationError, "没有可导入"):
                read_workbook(path)

    def test_template_has_guidance_validation_and_an_empty_input_table(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "template.xlsx"
            write_template(path)
            workbook = load_workbook(path, data_only=False)
            self.assertEqual(workbook.sheetnames, ["货物明细", "计算参数", "填写说明", "演示计算"])
            sheet = workbook["货物明细"]
            self.assertEqual(sheet.freeze_panes, "C2")
            self.assertIn("CargoInput", sheet.tables)
            self.assertEqual(tuple(sheet[1][i].value for i in range(8)), (
                "货物名称", "件数", "长", "宽", "高", "尺寸单位", "单件实际重量", "重量单位",
            ))
            self.assertIsNone(sheet[2][0].value)
            self.assertEqual(sheet[2][1].fill.fgColor.rgb[-6:], "FFF4D6")
            self.assertEqual(sheet[2][2].number_format, "0.########")
            validation_refs = {str(rule.sqref) for rule in sheet.data_validations.dataValidation}
            self.assertIn("B2:B5001", validation_refs)
            self.assertIn("F2:F5001", validation_refs)
            self.assertIn("H2:H5001", validation_refs)
            self.assertIn("货物明细", workbook["填写说明"]["A2"].value)
            params = workbook["计算参数"]
            self.assertEqual(len(params.data_validations.dataValidation), 3)
            demo = workbook["演示计算"]
            self.assertIn("合成数据", demo["A1"].value)
            self.assertEqual(demo["A9"].value, "体积重除数")
            self.assertEqual(demo["B12"].value, 0.196)
            self.assertEqual(demo["B13"].value, 32.2)
            self.assertAlmostEqual(demo["B14"].value, 32.666666666666664, places=12)
            self.assertEqual(demo["B15"].value, 39)
            self.assertEqual(demo["B16"].value, 33)
            self.assertEqual(demo["A18"].value, "不同渠道的除数、计费口径、进位及最低计费可能不同，请以该票承运渠道确认的书面规则为准。")
            workbook.close()

    def test_entered_template_round_trips_multiple_rows_and_custom_parameters(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "entered-template.xlsx"
            write_template(path)
            workbook = load_workbook(path)
            sheet = workbook["货物明细"]
            sheet.append(["Carton A", 3, 10, 20, 30, "in", 2.25, "lb"])
            sheet.append(["Carton B", 2, 50, 40, 30, "cm", 12.5, "kg"])
            params = workbook["计算参数"]
            params["B2"] = 5500
            params["B3"] = "逐件计算"
            params["B4"] = 0.25
            workbook.save(path)
            workbook.close()

            cargo_lines, settings = read_workbook(path)
            self.assertEqual(len(cargo_lines), 2)
            self.assertEqual(cargo_lines[0].pieces, 3)
            self.assertEqual(cargo_lines[0].dimension_unit, "in")
            self.assertEqual(cargo_lines[0].weight_unit, "lb")
            self.assertEqual(settings, CalculationSettings(
                Decimal("5500"), "per_piece", Decimal("0.25"),
            ))
            imported_result = calculate(cargo_lines, settings)
            self.assertEqual(imported_result.total_actual_weight_kg, Decimal("28.0617484975"))
            self.assertEqual(
                imported_result.selected_chargeable_weight_kg,
                imported_result.per_piece_chargeable_weight_kg,
            )

    def test_reordered_headers_and_chinese_unit_aliases_are_supported(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "source.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "货物明细"
            sheet.append(["重量单位", "单件实际重量", "尺寸单位", "高", "宽", "长", "件数", "货物名称"])
            sheet.append(["磅", 2, "英寸", 30, 20, 10, 1, "Box"])
            workbook.save(path)
            workbook.close()
            cargo_lines, settings = read_workbook(path)
        self.assertEqual(cargo_lines, [CargoLine(
            "Box", 1, Decimal("10"), Decimal("20"), Decimal("30"), "in",
            Decimal("2"), "lb",
        )])
        self.assertIsNone(settings)

    def test_import_validation_error_includes_excel_row(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "invalid-row.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "货物明细"
            sheet.append(["货物名称", "件数", "长", "宽", "高", "尺寸单位", "单件实际重量", "重量单位"])
            sheet.append(["Box", 1.5, 10, 10, 10, "cm", 1, "kg"])
            workbook.save(path)
            workbook.close()
            with self.assertRaisesRegex(FreightCalculationError, "第2行件数必须是整数"):
                read_workbook(path)

    def test_formula_like_text_is_exported_as_text(self):
        cargo = CargoLine(
            '=HYPERLINK("https://example.invalid")', 1, Decimal("1"), Decimal("1"),
            Decimal("1"), "cm", Decimal("1"), "kg",
        )
        settings = CalculationSettings()
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "safe-text.xlsx"
            write_results(path, [cargo], calculate([cargo], settings))
            workbook = load_workbook(path, data_only=False)
            cell = workbook["货物明细"]["A2"]
            self.assertEqual(cell.value, cargo.name)
            self.assertEqual(cell.data_type, "s")
            workbook.close()

    def test_exported_summary_matches_calculation_and_contains_no_formulas(self):
        lines = [
            CargoLine("Box A", 2, Decimal("25.4"), Decimal("50.8"), Decimal("76.2"),
                      "cm", Decimal("2.75"), "kg"),
            CargoLine("Box B", 1, Decimal("10"), Decimal("20"), Decimal("30"),
                      "in", Decimal("4"), "lb"),
        ]
        settings = CalculationSettings(Decimal("6000"), "shipment_total", Decimal("0.5"))
        result = calculate(lines, settings)
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "results.xlsx"
            write_results(path, lines, result)
            workbook = load_workbook(path, data_only=False)
            summary = workbook["汇总结果"]
            exported_metrics = [summary.cell(row, 2).value for row in range(2, 8)]
            expected_metrics = [
                float(result.total_volume_m3),
                float(result.total_actual_weight_kg),
                float(result.total_volumetric_weight_kg),
                float(result.per_piece_chargeable_weight_kg),
                float(result.shipment_total_chargeable_weight_kg),
                float(result.selected_chargeable_weight_kg),
            ]
            self.assertEqual(exported_metrics, expected_metrics)
            self.assertEqual(workbook["货物明细"].freeze_panes, "C2")
            self.assertEqual(workbook["货物明细"].auto_filter.ref, "A1:N3")
            formulas = [
                cell.coordinate
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if cell.data_type == "f"
            ]
            self.assertEqual(formulas, [])
            workbook.close()

    def test_rejects_non_xlsx(self):
        with self.assertRaisesRegex(FreightCalculationError, "仅支持"):
            read_workbook("not-an-excel-file.xls")


if __name__ == "__main__":
    unittest.main()
