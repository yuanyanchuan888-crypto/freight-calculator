"""Safe, schema-based XLSX import and export for the desktop calculator."""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from .calculator import CalculationResult, calculate
from .models import (
    MAX_PIECES_PER_LINE,
    CargoLine,
    CalculationSettings,
    FreightCalculationError,
)


IMPORT_HEADERS = (
    "货物名称",
    "件数",
    "长",
    "宽",
    "高",
    "尺寸单位",
    "单件实际重量",
    "重量单位",
)
MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_CARGO_ROWS = 5000


def _decimal_cell(value: object, label: str) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise FreightCalculationError(f"{label}不能为空，且必须是数字。")
    try:
        result = value if isinstance(value, Decimal) else Decimal(str(value).strip())
    except (InvalidOperation, ValueError, AttributeError):
        raise FreightCalculationError(f"{label}必须是有效数字。") from None
    if not result.is_finite():
        raise FreightCalculationError(f"{label}必须是有限数字。")
    return result


def _unit(value: object, label: str, aliases: dict[str, str]) -> str:
    key = str(value or "").strip().lower()
    if key not in aliases:
        allowed = "、".join(sorted(set(aliases.values())))
        raise FreightCalculationError(f"{label}只能填写：{allowed}。")
    return aliases[key]


def _read_settings(workbook) -> CalculationSettings | None:
    if "计算参数" not in workbook.sheetnames:
        return None
    values: dict[str, object] = {}
    for row in workbook["计算参数"].iter_rows(min_row=1, max_col=2, values_only=True):
        key = str(row[0] or "").strip()
        if key:
            values[key] = row[1]
    if "体积重除数" not in values:
        raise FreightCalculationError("Excel“计算参数”页缺少体积重除数。")
    divisor = _decimal_cell(values["体积重除数"], "体积重除数")
    mode_label = str(values.get("计费口径", "整票汇总")).strip()
    mode_map = {"逐件计算": "per_piece", "整票汇总": "shipment_total"}
    if mode_label not in mode_map:
        raise FreightCalculationError("计费口径应为“逐件计算”或“整票汇总”。")
    increment_value = values.get("进位单位(kg)")
    if increment_value is None or str(increment_value).strip() in ("", "不进位"):
        increment = None
    else:
        increment = _decimal_cell(increment_value, "进位单位(kg)")
    return CalculationSettings(
        volume_divisor=divisor,
        billing_mode=mode_map[mode_label],  # type: ignore[arg-type]
        rounding_increment_kg=increment,
    )


def read_workbook(path: str | Path) -> tuple[list[CargoLine], CalculationSettings | None]:
    """Read the documented XLSX template without executing formulas or macros."""
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise FreightCalculationError("仅支持 .xlsx 文件；请另存为 Excel 工作簿(.xlsx)。")
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise FreightCalculationError(f"无法读取文件：{exc}") from None
    if size > MAX_FILE_BYTES:
        raise FreightCalculationError("文件超过25 MB，请拆分后再导入。")

    try:
        from openpyxl import load_workbook
    except ImportError:
        raise FreightCalculationError(
            "Excel 功能需要 openpyxl。请在项目目录运行：python -m pip install -r requirements.txt"
        ) from None

    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:
        raise FreightCalculationError(f"无法打开 XLSX 文件：{exc}") from None

    try:
        if "货物明细" not in workbook.sheetnames:
            raise FreightCalculationError("工作簿缺少“货物明细”工作表。")
        sheet = workbook["货物明细"]
        if sheet.max_row > MAX_CARGO_ROWS + 1:
            raise FreightCalculationError(f"货物明细最多支持{MAX_CARGO_ROWS}行。")
        if sheet.max_column > 40:
            raise FreightCalculationError("货物明细列数异常，请使用本工具导出的模板。")
        rows = sheet.iter_rows(min_row=1, values_only=True)
        header_row = next(rows, None)
        if not header_row:
            raise FreightCalculationError("“货物明细”工作表为空。")
        headers = [str(value or "").strip() for value in header_row]
        if len(headers) != len(set(headers)):
            raise FreightCalculationError("Excel 表头存在重复列名。")
        missing = [header for header in IMPORT_HEADERS if header not in headers]
        if missing:
            raise FreightCalculationError("缺少必需列：" + "、".join(missing))
        indexes = {header: headers.index(header) for header in IMPORT_HEADERS}
        dimension_units = {
            "cm": "cm",
            "厘米": "cm",
            "in": "in",
            "英寸": "in",
        }
        weight_units = {
            "kg": "kg",
            "公斤": "kg",
            "千克": "kg",
            "lb": "lb",
            "lbs": "lb",
            "磅": "lb",
        }

        cargo_lines: list[CargoLine] = []
        for row_number, row in enumerate(rows, start=2):
            if not row or all(value is None or str(value).strip() == "" for value in row):
                continue
            if len(cargo_lines) >= MAX_CARGO_ROWS:
                raise FreightCalculationError(f"货物明细最多支持{MAX_CARGO_ROWS}行。")
            raw_name = row[indexes["货物名称"]]
            name = str(raw_name or "").strip()
            if not name:
                raise FreightCalculationError(f"第{row_number}行缺少货物名称。")
            pieces_value = _decimal_cell(row[indexes["件数"]], f"第{row_number}行件数")
            if pieces_value != pieces_value.to_integral_value():
                raise FreightCalculationError(f"第{row_number}行件数必须是整数。")
            try:
                cargo_lines.append(
                    CargoLine(
                        name=name,
                        pieces=int(pieces_value),
                        length=_decimal_cell(row[indexes["长"]], f"第{row_number}行长度"),
                        width=_decimal_cell(row[indexes["宽"]], f"第{row_number}行宽度"),
                        height=_decimal_cell(row[indexes["高"]], f"第{row_number}行高度"),
                        dimension_unit=_unit(
                            row[indexes["尺寸单位"]],
                            f"第{row_number}行尺寸单位",
                            dimension_units,
                        ),  # type: ignore[arg-type]
                        actual_weight_per_piece=_decimal_cell(
                            row[indexes["单件实际重量"]],
                            f"第{row_number}行单件实际重量",
                        ),
                        weight_unit=_unit(
                            row[indexes["重量单位"]],
                            f"第{row_number}行重量单位",
                            weight_units,
                        ),  # type: ignore[arg-type]
                    )
                )
            except FreightCalculationError as exc:
                raise FreightCalculationError(f"第{row_number}行：{exc}") from None
        if not cargo_lines:
            raise FreightCalculationError("工作簿中没有可导入的货物行。")
        return cargo_lines, _read_settings(workbook)
    finally:
        workbook.close()


def _style_header(sheet, row: int = 1) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment

    for cell in sheet[row]:
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173B57")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 28


def _write_text(cell, value: str) -> None:
    cell.value = str(value)
    cell.data_type = "s"


def _write_numeric(cell, value: Decimal) -> None:
    cell.value = float(value)
    cell.number_format = "0.########"


def _prepare_sheet(sheet) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    if sheet.freeze_panes is None:
        sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for column in sheet.columns:
        letter = column[0].column_letter
        max_length = min(max((len(str(c.value or "")) for c in column[:100]), default=12), 28)
        sheet.column_dimensions[letter].width = max(12, max_length + 2)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", color="263746")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F3F7FA")


def write_template(path: str | Path) -> None:
    """Create an empty, documented input template for .xlsx imports."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        raise FreightCalculationError(
            "Excel 功能需要 openpyxl。请在项目目录运行：python -m pip install -r requirements.txt"
        ) from None
    workbook = Workbook()
    workbook.properties.title = "国际物流计费计算资料模板"
    workbook.properties.subject = "体积重与计费重量计算输入模板"
    workbook.properties.creator = "International Freight Calculator"
    sheet = workbook.active
    sheet.title = "货物明细"
    sheet.append(list(IMPORT_HEADERS))
    sheet.append([None] * len(IMPORT_HEADERS))
    _style_header(sheet)
    sheet.freeze_panes = "C2"
    sheet.sheet_view.zoomScale = 90
    sheet.auto_filter.ref = "A1:H2"
    table = Table(displayName="CargoInput", ref="A1:H2")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)
    input_fill = PatternFill("solid", fgColor="FFF4D6")
    for cell in sheet[2]:
        cell.fill = input_fill
        cell.font = Font(name="Arial", color="263746")
        cell.alignment = Alignment(vertical="center")
    for column in (2, 3, 4, 5, 7):
        sheet.cell(2, column).number_format = "0.########"
    sheet.row_dimensions[2].height = 25

    pieces_validation = DataValidation(
        type="whole", operator="between", formula1="1", formula2=str(MAX_PIECES_PER_LINE),
        allow_blank=True,
    )
    pieces_validation.errorTitle = "件数无效"
    pieces_validation.error = "请输入 1 至 1,000,000 的整数。"
    pieces_validation.promptTitle = "填写件数"
    pieces_validation.prompt = "填写这一行相同规格的包装件数。"
    pieces_validation.showErrorMessage = True
    pieces_validation.showInputMessage = True
    sheet.add_data_validation(pieces_validation)
    pieces_validation.add("B2:B5001")

    positive_validation = DataValidation(
        type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
    )
    positive_validation.errorTitle = "数值无效"
    positive_validation.error = "请输入大于 0 的数字；最大值请以导入校验提示为准。"
    positive_validation.showErrorMessage = True
    sheet.add_data_validation(positive_validation)
    for column in ("C", "D", "E", "G"):
        positive_validation.add(f"{column}2:{column}5001")

    dimension_validation = DataValidation(type="list", formula1='"cm,in"', allow_blank=True)
    dimension_validation.errorTitle = "尺寸单位无效"
    dimension_validation.error = "请选择 cm 或 in。"
    dimension_validation.promptTitle = "尺寸单位"
    dimension_validation.prompt = "每行可单独选择厘米或英寸。"
    dimension_validation.showErrorMessage = True
    dimension_validation.showInputMessage = True
    sheet.add_data_validation(dimension_validation)
    dimension_validation.add("F2:F5001")

    weight_validation = DataValidation(type="list", formula1='"kg,lb"', allow_blank=True)
    weight_validation.errorTitle = "重量单位无效"
    weight_validation.error = "请选择 kg 或 lb。"
    weight_validation.promptTitle = "重量单位"
    weight_validation.prompt = "每行可单独选择公斤或磅。"
    weight_validation.showErrorMessage = True
    weight_validation.showInputMessage = True
    sheet.add_data_validation(weight_validation)
    weight_validation.add("H2:H5001")

    params = workbook.create_sheet("计算参数")
    params.append(["参数", "值", "说明"])
    _style_header(params)
    params.append(["体积重除数", 6000, "示例值；请按具体承运渠道规则确认。"])
    params.append(["计费口径", "整票汇总", "可选：逐件计算、整票汇总。"])
    params.append(["进位单位(kg)", None, "留空表示不进位；填入如 0.5 表示向上进位至0.5 kg。"])
    params.append(["单位填写", "货物明细", "尺寸单位每行填 cm 或 in；重量单位每行填 kg 或 lb。"])
    params.append(["重要说明", "本模板不代表任何货代或承运商的统一规则。", ""])
    divisor_validation = DataValidation(type="decimal", operator="greaterThan", formula1="0")
    divisor_validation.error = "体积重除数必须大于 0。"
    divisor_validation.showErrorMessage = True
    params.add_data_validation(divisor_validation)
    divisor_validation.add("B2")
    mode_validation = DataValidation(type="list", formula1='"逐件计算,整票汇总"')
    mode_validation.error = "请选择逐件计算或整票汇总。"
    mode_validation.showErrorMessage = True
    params.add_data_validation(mode_validation)
    mode_validation.add("B3")
    increment_validation = DataValidation(type="decimal", operator="greaterThan", formula1="0",
                                          allow_blank=True)
    increment_validation.errorTitle = "进位单位无效"
    increment_validation.error = "进位单位必须大于 0；留空表示不进位。"
    increment_validation.showErrorMessage = True
    params.add_data_validation(increment_validation)
    increment_validation.add("B4")
    _prepare_sheet(sheet)
    for column, width in {
        "A": 24, "B": 12, "C": 13, "D": 13, "E": 13, "F": 14,
        "G": 19, "H": 14,
    }.items():
        sheet.column_dimensions[column].width = width
    _prepare_sheet(params)
    params.column_dimensions["A"].width = 22
    params.column_dimensions["B"].width = 28
    params.column_dimensions["C"].width = 62
    for row_number in range(2, params.max_row + 1):
        params.row_dimensions[row_number].height = 30
    for cell in sheet[2]:
        cell.fill = input_fill

    guide = workbook.create_sheet("填写说明")
    guide.sheet_view.showGridLines = False
    guide.sheet_view.zoomScale = 95
    guide.append(["国际物流计费资料模板", "填写说明"])
    guide.merge_cells("A1:B1")
    guide["A1"].font = Font(name="Arial", size=14, bold=True, color="173B57")
    guide["A1"].alignment = Alignment(vertical="center")
    guide.row_dimensions[1].height = 30
    guide_rows = (
        ("1. 货物明细", "每行填写一种相同包装规格。件数表示该行包装的总件数。"),
        ("2. 尺寸与重量", "填写单件外箱长、宽、高和单件实际重量。单位可逐行选择 cm/in 与 kg/lb。"),
        ("3. 计算参数", "填写体积重除数、计费口径和进位单位。默认值仅作演示，可按具体渠道规则调整。"),
        ("4. 导入方式", "仅支持 .xlsx。导入公式不会执行，请将需要导入的输入值保存为普通单元格数据。"),
        ("5. 计费口径", "逐件计算与整票汇总可能产生不同结果。最低计费、计泡、分泡及进位规则须向承运渠道核实。"),
        ("6. 文件限制", "最多导入 5,000 行货物资料。输入页保持空白；独立“演示计算”页使用合成数据。"),
    )
    for row in guide_rows:
        guide.append(row)
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 96
    for row in guide.iter_rows(min_row=2, max_col=2):
        row[0].font = Font(name="Arial", bold=True, color="173B57")
        row[1].font = Font(name="Arial", color="263746")
        row[0].alignment = Alignment(vertical="top", wrap_text=True)
        row[1].alignment = Alignment(vertical="top", wrap_text=True)
        guide.row_dimensions[row[0].row].height = 32

    example = workbook.create_sheet("演示计算")
    example.sheet_view.showGridLines = False
    example.sheet_view.zoomScale = 90
    example.merge_cells("A1:L1")
    example["A1"] = "演示计算（全部为合成数据，不会从本页导入）"
    example["A1"].font = Font(name="Microsoft YaHei", size=14, bold=True, color="173B57")
    example["A1"].alignment = Alignment(vertical="center")
    example.row_dimensions[1].height = 30
    example.merge_cells("A2:L2")
    example["A2"] = "输入页“货物明细”仍为空白。以下仅用于核对计算口径，不代表任何承运商费率或实际货物。"
    example["A2"].font = Font(name="Microsoft YaHei", color="617384", italic=True)
    example["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    example.row_dimensions[2].height = 30

    demo_lines = [
        CargoLine("模拟纸箱 A", 3, Decimal("50"), Decimal("40"), Decimal("30"),
                  "cm", Decimal("8"), "kg"),
        CargoLine("模拟纸箱 B", 2, Decimal("20"), Decimal("20"), Decimal("20"),
                  "cm", Decimal("4.1"), "kg"),
    ]
    demo_settings = CalculationSettings(Decimal("6000"), "shipment_total", Decimal("0.5"))
    demo_result = calculate(demo_lines, demo_settings)
    example_headers = (
        "货物名称", "件数", "长", "宽", "高", "尺寸单位", "单件实重", "重量单位",
        "行体积(m³)", "行实重(kg)", "行体积重(kg)", "逐件口径计费重(kg)",
    )
    for column, heading in enumerate(example_headers, start=1):
        cell = example.cell(4, column, heading)
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173B57")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    example.row_dimensions[4].height = 34
    for row_number, calculation in enumerate(demo_result.lines, start=5):
        cargo = calculation.cargo
        values = (
            cargo.name, cargo.pieces, float(cargo.length), float(cargo.width), float(cargo.height),
            cargo.dimension_unit, float(cargo.actual_weight_per_piece), cargo.weight_unit,
            float(calculation.volume_m3), float(calculation.actual_weight_kg),
            float(calculation.volumetric_weight_kg), float(calculation.per_piece_chargeable_kg),
        )
        for column, value in enumerate(values, start=1):
            cell = example.cell(row_number, column, value)
            cell.font = Font(name="Microsoft YaHei", color="263746")
            cell.alignment = Alignment(vertical="center")
            if column in (2, 3, 4, 5, 7, 9, 10, 11, 12):
                cell.number_format = "0.########"
            if row_number % 2:
                cell.fill = PatternFill("solid", fgColor="F3F7FA")
    example.freeze_panes = "B5"

    example["A8"] = "示例参数与汇总"
    example["A8"].font = Font(name="Microsoft YaHei", bold=True, color="173B57")
    example_rows = (
        ("体积重除数", float(demo_settings.volume_divisor), "cm³/kg，仅为演示参数"),
        ("计费口径", "整票汇总", "先汇总实重与体积重，再比较并进位"),
        ("进位单位", float(demo_settings.rounding_increment_kg), "kg，按增量向上进位"),
        ("总体积", float(demo_result.total_volume_m3), "m³"),
        ("总实际重量", float(demo_result.total_actual_weight_kg), "kg"),
        ("总体积重", float(demo_result.total_volumetric_weight_kg), "kg"),
        ("逐件口径计费重量", float(demo_result.per_piece_chargeable_weight_kg), "kg"),
        ("整票汇总计费重量", float(demo_result.shipment_total_chargeable_weight_kg), "kg"),
    )
    for row_number, row in enumerate(example_rows, start=9):
        for column, value in enumerate(row, start=1):
            example.cell(row_number, column, value)
        example.cell(row_number, 2).number_format = "0.########"
    example.merge_cells("A18:L18")
    example["A18"] = "不同渠道的除数、计费口径、进位及最低计费可能不同，请以该票承运渠道确认的书面规则为准。"
    example["A18"].font = Font(name="Microsoft YaHei", color="8A4B08", italic=True)
    example["A18"].alignment = Alignment(wrap_text=True, vertical="center")
    example.row_dimensions[18].height = 30
    example_widths = (20, 9, 9, 9, 9, 12, 12, 12, 14, 14, 15, 20)
    for index, width in enumerate(example_widths, start=1):
        example.column_dimensions[chr(64 + index)].width = width
    workbook.save(path)
    workbook.close()


def write_results(
    path: str | Path,
    cargo_lines: Iterable[CargoLine],
    result: CalculationResult,
) -> None:
    try:
        from openpyxl import Workbook
    except ImportError:
        raise FreightCalculationError(
            "Excel 功能需要 openpyxl。请在项目目录运行：python -m pip install -r requirements.txt"
        ) from None

    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总结果"
    summary.append(["指标", "数值", "单位或说明"])
    _style_header(summary)
    metrics = [
        ("总体积", result.total_volume_m3, "m³"),
        ("总实际重量", result.total_actual_weight_kg, "kg"),
        ("总体积重", result.total_volumetric_weight_kg, "kg"),
        ("逐件口径计费重量", result.per_piece_chargeable_weight_kg, "kg"),
        ("整票汇总计费重量", result.shipment_total_chargeable_weight_kg, "kg"),
        ("当前选中计费重量", result.selected_chargeable_weight_kg, "kg"),
    ]
    for label, value, unit in metrics:
        summary.append([label, float(value), unit])
        summary.cell(summary.max_row, 2).number_format = "0.########"
    summary.append(["体积重除数", float(result.settings.volume_divisor), "cm³/kg 示例参数"])
    summary.cell(summary.max_row, 2).number_format = "0.####"
    mode = "逐件计算" if result.settings.billing_mode == "per_piece" else "整票汇总"
    summary.append(["当前计费口径", mode, "计费及进位规则须以具体渠道书面确认为准"])
    increment = (
        "不进位"
        if result.settings.rounding_increment_kg is None
        else float(result.settings.rounding_increment_kg)
    )
    summary.append(["进位单位(kg)", increment, "按设置向上进位"])
    summary.append(["核算说明", "仅为计算示例，不代表承运商实际账单。", ""])
    _prepare_sheet(summary)
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 24
    summary.column_dimensions["C"].width = 58

    detail = workbook.create_sheet("货物明细")
    detail_headers = [
        *IMPORT_HEADERS,
        "行体积(m³)",
        "单件实际重量(kg)",
        "行实际重量(kg)",
        "单件体积重(kg)",
        "行体积重(kg)",
        "逐件口径计费重量(kg)",
    ]
    detail.append(detail_headers)
    _style_header(detail)
    detail.freeze_panes = "C2"
    detail.sheet_view.zoomScale = 85
    for item in result.lines:
        line = item.cargo
        row = [
            line.name,
            line.pieces,
            float(line.length),
            float(line.width),
            float(line.height),
            line.dimension_unit,
            float(line.actual_weight_per_piece),
            line.weight_unit,
            float(item.volume_m3),
            float(item.actual_weight_kg_per_piece),
            float(item.actual_weight_kg),
            float(item.volumetric_weight_kg_per_piece),
            float(item.volumetric_weight_kg),
            float(item.per_piece_chargeable_kg),
        ]
        detail.append(row)
        _write_text(detail.cell(detail.max_row, 1), line.name)
        for cell in detail[detail.max_row]:
            if cell.column not in (1, 2, 6, 8):
                cell.number_format = "0.########"
    _prepare_sheet(detail)
    for column, width in {
        "A": 24, "B": 10, "C": 12, "D": 12, "E": 12, "F": 14,
        "G": 20, "H": 14, "I": 16, "J": 21, "K": 18, "L": 20,
        "M": 18, "N": 22,
    }.items():
        detail.column_dimensions[column].width = width
    detail.auto_filter.ref = detail.dimensions

    params = workbook.create_sheet("计算参数")
    params.append(["参数", "值", "说明"])
    _style_header(params)
    params.append(["体积重除数", float(result.settings.volume_divisor), "示例参数，按具体渠道确认。"])
    params.append(["计费口径", mode, "逐件计算或整票汇总。"])
    params.append(["进位单位(kg)", increment, "空白/不进位表示不作进位。"])
    params.append(["生成时间", datetime.now().astimezone().isoformat(timespec="seconds"), "本地时间"])
    params.append(["规则说明", "本工具不内置或代表任何一家承运商的计费规则。", ""])
    _prepare_sheet(params)
    params.column_dimensions["A"].width = 20
    params.column_dimensions["B"].width = 48
    params.column_dimensions["C"].width = 58
    workbook.save(path)
    workbook.close()
